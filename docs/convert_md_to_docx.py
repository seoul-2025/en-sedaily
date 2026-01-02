#!/usr/bin/env python3
"""
Markdown to Word Document Converter
Converts TECHNICAL_REPORT_SEO_AEO.md to .docx format
"""

import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def parse_markdown_to_docx(md_file, output_file):
    """Convert markdown file to Word document with proper formatting"""

    # Read markdown file
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()

    # Create Word document
    doc = Document()

    # Set document margins (1 inch all around)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Process markdown line by line
    lines = content.split('\n')
    i = 0

    while i < len(lines):
        line = lines[i]

        # Skip empty lines at document start
        if not line.strip() and not doc.paragraphs:
            i += 1
            continue

        # H1 - Main Title
        if line.startswith('# ') and not line.startswith('## '):
            text = line[2:].strip()
            p = doc.add_heading(text, level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.runs[0]
            run.font.size = Pt(24)
            run.font.color.rgb = RGBColor(0, 51, 102)  # Dark blue
            run.font.bold = True

        # H2 - Section Headers
        elif line.startswith('## '):
            text = line[3:].strip()
            p = doc.add_heading(text, level=2)
            run = p.runs[0]
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(0, 102, 204)  # Blue

        # H3 - Subsection Headers
        elif line.startswith('### '):
            text = line[4:].strip()
            p = doc.add_heading(text, level=3)
            run = p.runs[0]
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(51, 51, 51)  # Dark gray

        # H4 - Minor Headers
        elif line.startswith('#### '):
            text = line[5:].strip()
            p = doc.add_heading(text, level=4)
            run = p.runs[0]
            run.font.size = Pt(12)
            run.font.bold = True

        # Horizontal Rule
        elif line.strip() == '---':
            p = doc.add_paragraph()
            p.add_run('_' * 80)
            run = p.runs[0]
            run.font.color.rgb = RGBColor(200, 200, 200)

        # Code blocks
        elif line.strip().startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1

            # Add code block
            p = doc.add_paragraph('\n'.join(code_lines))
            p.style = 'Normal'
            run = p.runs[0]
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)

            # Gray background simulation
            run.font.color.rgb = RGBColor(60, 60, 60)

        # Bullet list (unordered)
        elif line.strip().startswith('- ') or line.strip().startswith('* '):
            text = re.sub(r'^[\-\*]\s+', '', line).strip()
            text = format_inline_markdown(text)
            p = doc.add_paragraph(text, style='List Bullet')
            p.paragraph_format.left_indent = Inches(0.25)

        # Numbered list
        elif re.match(r'^\d+\.\s+', line.strip()):
            text = re.sub(r'^\d+\.\s+', '', line).strip()
            text = format_inline_markdown(text)
            p = doc.add_paragraph(text, style='List Number')
            p.paragraph_format.left_indent = Inches(0.25)

        # Tables
        elif '|' in line and line.strip().startswith('|'):
            table_lines = []
            # Collect all table lines
            while i < len(lines) and '|' in lines[i]:
                table_lines.append(lines[i])
                i += 1
            i -= 1  # Back one line

            # Parse table
            if len(table_lines) > 2:  # Header + separator + at least one row
                # Remove separator line (second line)
                header = table_lines[0]
                rows = table_lines[2:]

                # Parse columns
                headers = [h.strip() for h in header.split('|') if h.strip()]

                # Create table
                table = doc.add_table(rows=1 + len(rows), cols=len(headers))
                table.style = 'Light Grid Accent 1'

                # Header row
                for col_idx, header_text in enumerate(headers):
                    cell = table.rows[0].cells[col_idx]
                    cell.text = header_text
                    # Bold header
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.bold = True
                            run.font.size = Pt(10)

                # Data rows
                for row_idx, row_line in enumerate(rows, start=1):
                    cells_data = [c.strip() for c in row_line.split('|') if c.strip()]
                    for col_idx, cell_text in enumerate(cells_data):
                        if col_idx < len(headers):
                            cell = table.rows[row_idx].cells[col_idx]
                            cell.text = cell_text
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    run.font.size = Pt(9)

                # Add space after table
                doc.add_paragraph()

        # Blockquote
        elif line.strip().startswith('> '):
            text = line.strip()[2:]
            p = doc.add_paragraph(text)
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.right_indent = Inches(0.5)
            run = p.runs[0]
            run.font.italic = True
            run.font.color.rgb = RGBColor(100, 100, 100)

        # Regular paragraph
        elif line.strip():
            text = format_inline_markdown(line.strip())
            if text:
                p = doc.add_paragraph()
                add_formatted_text(p, text)

        # Empty line
        else:
            if doc.paragraphs:  # Only add spacing if not at document start
                doc.add_paragraph()

        i += 1

    # Save document
    doc.save(output_file)
    print(f"✅ Document saved: {output_file}")

def format_inline_markdown(text):
    """Remove markdown formatting for inline text"""
    # Bold: **text** or __text__
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'__(.+?)__', r'\1', text)

    # Italic: *text* or _text_
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'_(.+?)_', r'\1', text)

    # Inline code: `code`
    text = re.sub(r'`(.+?)`', r'\1', text)

    # Links: [text](url)
    text = re.sub(r'\[(.+?)\]\(.+?\)', r'\1', text)

    # Remove emojis (optional - they work in Word but might render poorly)
    # text = re.sub(r'[\U0001F300-\U0001F9FF]', '', text)

    return text

def add_formatted_text(paragraph, text):
    """Add text to paragraph with inline formatting"""
    # Handle bold (**text**)
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.font.bold = True
        elif part:
            paragraph.add_run(part)

if __name__ == "__main__":
    input_file = "TECHNICAL_REPORT_SEO_AEO.md"
    output_file = "TECHNICAL_REPORT_SEO_AEO.docx"

    print(f"📄 Converting {input_file} to Word document...")
    parse_markdown_to_docx(input_file, output_file)
    print(f"✅ Conversion complete!")
    print(f"📁 Output: {output_file}")
